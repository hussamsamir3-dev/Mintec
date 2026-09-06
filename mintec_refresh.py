#!/usr/bin/env python3
"""
mintec_refresh.py — pull live prices from Mintec and rebuild the Market Desk.

Built against Mintec API technical documentation v2.4.

  Auth      OAuth 2.0 client_credentials -> bearer token, valid 1 hour
            POST https://identity.mintecanalytics.com/connect/token
            scope: export_api   (the Forecast API reuses this same scope)
  Base      https://public-api.mintecanalytics.com
  Actuals   POST /v2/export/series/mintec/points     up to 100 codes per call
  Forecast  POST /v2/export/series/forecast/points   same shape, seriesType=forecast
  Dates     DD/MM/YYYY, in and out

The material list can be either Book1.xlsx or a plain materials.csv holding
one "Name, MintecCode" per line. With --from-api the CSV is enough: currency,
unit, frequency, category and description all come back from Mintec.

Commands
────────
  python mintec_refresh.py --probe
      Gets a token, reports your subscription size, and checks a sample code
      against both the price and forecast endpoints.

  python mintec_refresh.py --from-api
      Pulls actuals and forecasts for every code in the workbook, merges them,
      rewrites the workbook and rebuilds the dashboard.

  python mintec_refresh.py --from-api --every 60
      The same, on a loop, every 60 minutes, forever. Leave it running on the
      machine driving the screen and the dashboard stays current on its own.
      Ctrl-C stops it. Survives transient failures — a bad hour is logged and
      the next hour is tried as normal.

  python mintec_refresh.py --from-workbook Book1.xlsx
      Rebuilds the dashboard from a workbook refreshed by the Excel add-in.
      No credentials needed.

  python mintec_refresh.py --catalogue "milk powder"
      Searches every series your subscription can see. Use it to find the
      Mintec code for a material you want to add to the dashboard.

Credentials — Mintec Analytics > User profile > APIs Access
  export MINTEC_CLIENT_ID='...'
  export MINTEC_CLIENT_SECRET='...'

Never put credentials in this file.

Requires: openpyxl, requests   ->   pip install openpyxl requests
"""

import argparse
import datetime as dt
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is missing. Run: pip install openpyxl requests")

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_WB = os.path.join(HERE, "Book1.xlsx")
DEFAULT_HTML = os.path.join(HERE, "Materials_Market_Desk.html")

AUTHORITY = os.environ.get("MINTEC_AUTHORITY", "https://identity.mintecanalytics.com/connect/token")
BASE = os.environ.get("MINTEC_BASE", "https://public-api.mintecanalytics.com").rstrip("/")
SCOPE = os.environ.get("MINTEC_SCOPE", "")   # blank = try the known variants in turn

# The v2.4 doc lists the scope as "export_api, import_api", which is ambiguous:
# OAuth separates scopes with spaces, but some servers want a single value or a
# comma-separated pair. Rather than guess, try each and keep whichever works.
SCOPE_CANDIDATES = ["export_api",
                    "export_api import_api",
                    "export_api,import_api",
                    "import_api"]

HIST_YEARS = int(os.environ.get("MINTEC_HISTORY_YEARS", "3"))
FWD_YEARS = int(os.environ.get("MINTEC_FORWARD_YEARS", "3"))
BATCH = 100                      # doc: up to 100 series per POST, no paging
SKIP_TABS = {"DataValidationListsDONOTDELETE", "commo"}
WB_DATE = re.compile(r"^\d{2}/\d{2}/\d{4}$")


# ═══════════════════════════ auth ═══════════════════════════
_TOKEN = [None, dt.datetime.min, None]


def get_token(session, force=False):
    if _TOKEN[0] and not force and _TOKEN[1] > dt.datetime.now():
        return _TOKEN[0]
    cid = os.environ.get("MINTEC_CLIENT_ID")
    secret = os.environ.get("MINTEC_CLIENT_SECRET")
    if not cid or not secret:
        sys.exit("Set MINTEC_CLIENT_ID and MINTEC_CLIENT_SECRET.\n"
                 "Both are on Mintec Analytics > User profile > APIs Access.")
    tried = []
    for scope in ([SCOPE] if SCOPE else SCOPE_CANDIDATES):
        r = session.post(AUTHORITY,
                         data={"grant_type": "client_credentials", "client_id": cid,
                               "client_secret": secret, "scope": scope},
                         headers={"Content-Type": "application/x-www-form-urlencoded"},
                         timeout=45)
        if r.status_code == 200:
            if scope != SCOPE_CANDIDATES[0]:
                print(f"  (scope '{scope}' accepted)")
            _TOKEN[2] = scope
            break
        tried.append(f"{scope!r} -> {r.status_code} {r.text[:90]}")
    else:
        detail = "\n    ".join(tried)
        raise RuntimeError(
            "no scope was accepted. Mintec replied:\n    " + detail +
            "\n  If every line says invalid_scope, the credential is valid but this"
            "\n  account is not enabled for the Export API — ask Mintec support to"
            "\n  add export_api to your subscription.")
    body = r.json()
    tok = body.get("access_token")
    if not tok:
        raise RuntimeError(f"no access_token in response: {r.text[:200]}")
    ttl = int(body.get("expires_in") or 3600)
    _TOKEN[0] = tok
    _TOKEN[1] = dt.datetime.now() + dt.timedelta(seconds=max(ttl - 120, 60))
    return tok


def call(session, method, path, params=None, body=None, retry=True):
    r = session.request(
        method, BASE + path,
        headers={"Authorization": "Bearer " + get_token(session),
                 "Accept": "application/json",
                 "Content-Type": "application/json"},
        params=params or {}, json=body, timeout=120)
    if r.status_code == 401 and retry:
        get_token(session, force=True)
        return call(session, method, path, params, body, retry=False)
    return r


# ═══════════════════════ response shapes ════════════════════
def content_of(r):
    """Every v2 response is {"content": ..., "code": n, "messages": [...]}."""
    if r.status_code != 200:
        raise RuntimeError(f"HTTP {r.status_code}: {r.text[:160]}")
    body = r.json()
    for m in body.get("messages") or []:
        if isinstance(m, dict) and m.get("message"):
            print(f"      note [{m.get('key')}] {m['message']}")
    return body.get("content")


def to_iso(d):
    """Mintec sends DD/MM/YYYY."""
    s = str(d)[:10]
    m = re.match(r"^(\d{2})[/-](\d{2})[/-](\d{4})$", s)
    if m:
        return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    try:
        dt.date.fromisoformat(s)
        return s
    except ValueError:
        return None


def from_iso(s):
    y, m, d = s.split("-")
    return f"{d}/{m}/{y}"


def series_from(obj):
    """One series object -> (metadata dict, sorted [[iso, value]])."""
    pts = {}
    for p in obj.get("points") or []:
        if not isinstance(p, dict):
            continue
        d, v = to_iso(p.get("date")), p.get("value")
        if d is None or v is None:
            continue
        try:
            pts[d] = round(float(v), 4)
        except (TypeError, ValueError):
            continue
    meta = {
        "code": obj.get("seriesCode"),
        "name": obj.get("seriesName"),
        "description": obj.get("seriesName"),
        "specification": obj.get("seriesDescription"),
        "currency": obj.get("currencyName"),
        "unit": obj.get("unitName"),
        "frequency": obj.get("frequencyName"),
        "seriesType": obj.get("seriesType"),
        "category": obj.get("category"),
        "subCategory": obj.get("subCategory"),
        "origin": obj.get("countryOfOriginName"),
        "delivery": obj.get("countryOfDeliveryName"),
        "similarity": obj.get("statisticalSimilarity"),
    }
    return meta, [[d, v] for d, v in sorted(pts.items())]


# ═══════════════════════ API methods ════════════════════════
def fetch_points(session, series_type, codes, start, end, cat=True):
    """
    POST /v2/export/series/{seriesType}/points
    Up to 100 codes per request, no pagination on the result.
    """
    out = {}
    for i in range(0, len(codes), BATCH):
        chunk = codes[i:i + BATCH]
        params = {}
        if cat:
            params["catMetadata"] = "true"
        if series_type == "forecast":
            params["seriesType"] = "true"
        r = call(session, "POST", f"/v2/export/series/{series_type}/points", params,
                 {"seriesCodes": chunk, "startDate": from_iso(start), "endDate": from_iso(end)})
        content = content_of(r)
        if content is None:
            continue
        if isinstance(content, dict):
            content = [content]
        for obj in content:
            meta, pts = series_from(obj)
            if meta["code"]:
                out[meta["code"].strip().upper()] = (meta, pts)
    return out


def list_catalogue(session, series_type="mintec", page_size=1000):
    """GET /v2/export/series/{seriesType} — every series in your subscription."""
    items, page = [], 0
    while True:
        r = call(session, "GET", f"/v2/export/series/{series_type}",
                 {"PageIndex": page, "PageSize": page_size, "catMetadata": "true"})
        content = content_of(r)
        if content is None:
            break
        batch = content.get("items") if isinstance(content, dict) else content
        if not batch:
            break
        items.extend(batch)
        total = content.get("totalItems") if isinstance(content, dict) else None
        if total is None or len(items) >= total or len(batch) < page_size:
            break
        page += 1
    return items


# ═══════════════════════ merge ══════════════════════════════
def merge(actual_pts, forecast_pts, today=None):
    """
    Settled prices win. Forecast fills only the dates beyond the last
    settled observation, so the dashboard's actual/forecast boundary is real
    rather than assumed.
    """
    today = today or dt.date.today().isoformat()
    settled = [p for p in actual_pts if p[0] <= today]
    if not settled:
        settled = list(actual_pts)
    cutoff = settled[-1][0] if settled else today
    forward = [p for p in forecast_pts if p[0] > cutoff]
    return settled + forward, cutoff


# ═══════════════════════ workbook ═══════════════════════════
def read_list(path, quiet=False):
    """
    A plain-text material list: one line per material, "Name, MintecCode".
    Blank lines and lines starting with # are ignored. Everything else about
    each series — currency, unit, frequency, category — comes from the API.
    """
    out = {}
    for raw in open(path, encoding="utf-8-sig"):
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        parts = [p.strip() for p in line.split(",")]
        if len(parts) < 2 or not parts[0] or not parts[1]:
            continue
        if parts[0].lower() in ("name", "material"):      # header row
            continue
        name, code = parts[0], parts[1]
        out[name] = {"name": name, "code": code, "category": "Other", "series": []}
        if not quiet:
            print(f"  · {name:<11} {code}")
    return out


def read_source(path, quiet=False):
    """Accept either the Excel workbook or the plain material list."""
    if path.lower().endswith((".csv", ".txt", ".list")):
        return read_list(path, quiet)
    return read_workbook(path, quiet)


def read_workbook(path, quiet=False):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for name in wb.sheetnames:
        if name in SKIP_TABS:
            continue
        ws = wb[name]
        meta, series, code = {}, [], None
        for i, row in enumerate(ws.iter_rows(max_col=4, values_only=True), 1):
            a, b, c, d = (list(row) + [None] * 4)[:4]
            if i == 2 and d:
                code = str(d).strip()
            if isinstance(c, str) and d is not None:
                key = c.strip()
                if WB_DATE.match(key):
                    if isinstance(d, (int, float)):
                        series.append([to_iso(key), round(float(d), 4)])
                elif key:
                    meta[key] = d.date().isoformat() if isinstance(d, dt.datetime) else d
            if isinstance(a, str) and b is not None:
                meta.setdefault(a.strip().rstrip("*"),
                                b.date().isoformat() if isinstance(b, dt.datetime) else b)
        if not series:
            if not quiet:
                print(f"  · {name}: no price rows, skipped")
            continue
        series.sort()
        out[name] = {
            "name": name, "code": code,
            "description": meta.get("Description"), "specification": meta.get("Specification"),
            "frequency": meta.get("Frequency"), "currency": meta.get("Currency"),
            "unit": meta.get("Unit"), "seriesType": meta.get("Series Type"),
            "category": meta.get("Category") or "Other", "subCategory": meta.get("Sub-Category"),
            "origin": meta.get("Origin"), "delivery": meta.get("Delivery"),
            "series": series,
        }
        if not quiet:
            print(f"  · {name:<11} {str(code or '-'):<7} {len(series):>5} pts  "
                  f"{series[0][0]} -> {series[-1][0]}")
    wb.close()
    return out


def write_workbook(materials, path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, m in materials.items():
        ws = wb.create_sheet(name[:31])
        ws["D1"] = "Mintec codes go here"
        ws["D2"] = m.get("code") or ""
        r = 3
        for k in ("description", "specification", "frequency", "currency", "unit",
                  "seriesType", "category", "subCategory", "origin", "delivery"):
            label = {"seriesType": "Series Type", "subCategory": "Sub-Category"}.get(k, k.capitalize())
            ws.cell(r, 3, label)
            ws.cell(r, 4, m.get(k))
            r += 1
        for d, v in m["series"]:
            ws.cell(r, 3, from_iso(d))
            ws.cell(r, 4, v)
            r += 1
    wb.save(path)
    print(f"Workbook written: {path}")


# ═══════════════════════ commands ═══════════════════════════
def _session():
    try:
        import requests
    except ImportError:
        sys.exit("requests is missing. Run: pip install requests")
    return requests.Session()


def probe(materials):
    session = _session()
    sample = next((m["code"].strip() for m in materials.values() if (m.get("code") or "").strip()), "BCRD")
    print(f"Base   {BASE}\nScope  {SCOPE or ' / '.join(SCOPE_CANDIDATES)}\nSample {sample}\n")

    try:
        get_token(session)
        print(f"  token issued (scope: {_TOKEN[2]})")
    except Exception as exc:                                        # noqa: BLE001
        print(f"  token refused: {exc}")
        return

    for label, path in (("currencies", "/v2/currencies"), ("units", "/v2/units"),
                        ("frequencies", "/v2/frequencies")):
        r = call(session, "GET", path)
        n = len(content_of(r) or []) if r.status_code == 200 else 0
        print(f"  {label:<12} {r.status_code}  {n} entries")

    for stype in ("mintec", "forecast"):
        r = call(session, "GET", f"/v2/export/series/{stype}", {"PageSize": 1, "catMetadata": "true"})
        if r.status_code == 200:
            c = content_of(r) or {}
            total = c.get("totalItems") if isinstance(c, dict) else "?"
            print(f"  {stype:<12} 200  {total} series in your subscription")
        else:
            extra = "  (Forecast API is a separate subscription)" if stype == "forecast" else ""
            print(f"  {stype:<12} {r.status_code}{extra}")

    today = dt.date.today()
    start = (today - dt.timedelta(days=365)).isoformat()
    end = (today + dt.timedelta(days=730)).isoformat()
    for stype in ("mintec", "forecast"):
        try:
            got = fetch_points(session, stype, [sample], start, end)
            meta, pts = got.get(sample.upper(), ({}, []))
            print(f"  {stype:<12} {sample}: {len(pts)} points"
                  + (f"  {pts[0][0]} -> {pts[-1][0]}  {meta.get('currency')}/{meta.get('unit')}" if pts else ""))
        except Exception as exc:                                    # noqa: BLE001
            print(f"  {stype:<12} {sample}: {exc}")


def catalogue(term):
    session = _session()
    get_token(session)
    term = term.lower()
    for stype in ("mintec", "forecast"):
        try:
            items = list_catalogue(session, stype)
        except Exception as exc:                                    # noqa: BLE001
            print(f"{stype}: {exc}")
            continue
        hits = [i for i in items
                if term in " ".join(str(i.get(k) or "") for k in
                                    ("seriesCode", "seriesName", "seriesDescription",
                                     "category", "subCategory")).lower()]
        print(f"\n{stype}: {len(hits)} of {len(items)} series match '{term}'")
        for i in hits[:40]:
            print(f"  {str(i.get('seriesCode')):<8} {str(i.get('seriesName'))[:46]:<46} "
                  f"{str(i.get('currencyName') or '')[:12]:<12} {str(i.get('unitName') or '')[:14]:<14} "
                  f"{i.get('frequencyName') or ''}")
        if len(hits) > 40:
            print(f"  ... {len(hits) - 40} more")


def refresh_from_api(materials):
    session = _session()
    get_token(session)
    today = dt.date.today()
    start = (today - dt.timedelta(days=365 * HIST_YEARS)).isoformat()
    end = (today + dt.timedelta(days=365 * FWD_YEARS)).isoformat()

    wanted = {}
    for name, m in materials.items():
        c = (m.get("code") or "").strip().upper()
        if c:
            wanted[c] = name
        else:
            print(f"  x {name}: no Mintec code in the workbook")
    if not wanted:
        sys.exit("No Mintec codes found in the workbook.")
    codes = list(wanted)

    print(f"Requesting {len(codes)} codes, {start} to {end}")
    print("  settled prices ...")
    actual = fetch_points(session, "mintec", codes, start, end)
    print(f"    {len(actual)} series returned")

    forecast = {}
    try:
        print("  forward curves ...")
        forecast = fetch_points(session, "forecast", codes, start, end)
        print(f"    {len(forecast)} series returned")
    except Exception as exc:                                        # noqa: BLE001
        print(f"    forecast unavailable ({exc}); continuing with settled prices only")

    ok = 0
    for code, name in wanted.items():
        a_meta, a_pts = actual.get(code, ({}, []))
        f_meta, f_pts = forecast.get(code, ({}, []))
        if not a_pts and not f_pts:
            print(f"  x  {name:<11} {code:<7} nothing returned")
            continue
        pts, cutoff = merge(a_pts, f_pts)
        m = materials[name]
        m["series"] = pts
        for k in ("description", "specification", "currency", "unit", "frequency",
                  "category", "subCategory", "origin", "delivery"):
            v = a_meta.get(k) or f_meta.get(k)
            if v:
                m[k] = v
        if f_meta.get("seriesType"):
            m["seriesType"] = f_meta["seriesType"]
        ok += 1
        print(f"  ok {name:<11} {code:<7} {len(a_pts):>4} settled to {cutoff}"
              f"  + {len([p for p in pts if p[0] > cutoff]):>4} forward")
    print(f"\n{ok} of {len(wanted)} materials refreshed.")
    if not ok:
        sys.exit("Nothing came back. Run --probe to check the subscription.")
    return materials


def rebuild_html(materials, html_path, source_label="workbook"):
    if not os.path.exists(html_path):
        sys.exit(f"Can't find {html_path}. Keep it next to this script.")
    html = open(html_path, encoding="utf-8").read()
    payload = json.dumps(materials, separators=(",", ":"))
    new, n = re.subn(r"window\.__MATERIALS__\s*=\s*\{.*?\};",
                     lambda _: "window.__MATERIALS__ = " + payload + ";",
                     html, count=1, flags=re.S)
    if not n:
        sys.exit("Couldn't find the embedded data block in the HTML.")
    stamp = 'window.__BUILT__ = "%s, %s";' % (source_label, dt.datetime.now().strftime("%d %b %Y %H:%M"))
    new = re.sub(r'window\.__BUILT__\s*=\s*"[^"]*";', lambda _: stamp, new, count=1)
    open(html_path, "w", encoding="utf-8").write(new)
    pts = sum(len(m["series"]) for m in materials.values())
    newest = max(p[0] for m in materials.values() for p in m["series"])
    print(f"Dashboard rebuilt: {html_path}")
    print(f"  {len(materials)} materials · {pts:,} observations · curves to {newest}")


def main():
    ap = argparse.ArgumentParser(description="Refresh the Direct Materials Market Desk.")
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--probe", action="store_true", help="check credentials and subscription")
    g.add_argument("--from-api", action="store_true", help="pull live prices and forecasts")
    g.add_argument("--from-workbook", metavar="XLSX", nargs="?", const=DEFAULT_WB,
                   help="rebuild from a workbook you already refreshed")
    g.add_argument("--catalogue", metavar="TERM", help="search your subscription for a series code")
    ap.add_argument("--workbook", default=DEFAULT_WB,
                    help="material list: Book1.xlsx, or a materials.csv of 'Name, Code' lines")
    ap.add_argument("--html", default=DEFAULT_HTML)
    ap.add_argument("--out-workbook")
    ap.add_argument("--every", type=int, metavar="MINUTES",
                    help="repeat the refresh on a loop, e.g. --every 60")
    args = ap.parse_args()

    if args.catalogue:
        catalogue(args.catalogue)
        return

    wb_path = args.from_workbook if args.from_workbook else args.workbook
    materials = read_source(wb_path, quiet=args.probe)
    if not materials:
        sys.exit("No material tabs found in the workbook.")

    if args.probe:
        probe(materials)
        return

    def once():
        mats = read_source(wb_path, quiet=True) if args.every else materials
        if args.from_api:
            mats = refresh_from_api(mats)
            if not wb_path.lower().endswith((".csv", ".txt", ".list")):
                write_workbook(mats, args.out_workbook or wb_path)
            elif args.out_workbook:
                write_workbook(mats, args.out_workbook)
        rebuild_html(mats, args.html,
                     "Mintec API" if args.from_api else f"workbook {os.path.basename(wb_path)}")

    if not args.every:
        once()
        print("Done.")
        return

    import time
    print(f"Refreshing every {args.every} minutes. Ctrl-C to stop.\n")
    while True:
        stamp = dt.datetime.now().strftime("%d %b %H:%M")
        print(f"───── {stamp} ─────")
        try:
            once()
        except SystemExit as exc:
            print(f"  stopped: {exc}")
        except Exception as exc:                                    # noqa: BLE001
            print(f"  refresh failed: {exc}")
            print("  the dashboard keeps its last good data; trying again next cycle")
        nxt = dt.datetime.now() + dt.timedelta(minutes=args.every)
        print(f"  next refresh {nxt.strftime('%H:%M')}\n")
        try:
            time.sleep(args.every * 60)
        except KeyboardInterrupt:
            print("Stopped.")
            return


if __name__ == "__main__":
    main()
